import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Auto Pivot Stock SPR JABO", layout="wide")

st.title("📦 Aplikasi Auto-Pivot Stock (2 Sheet: Device & Acc)")
st.write("Sistem otomatis membagi data menjadi Sheet Device & Accessories berdasarkan file Master Kategori dengan format presisi SPR JABO.")

# 1. Fitur Upload 2 File (Mentahan & Master)
col1, col2 = st.columns(2)
with col1:
    uploaded_raw = st.file_uploader("1️⃣ Upload File MENTAHAN (.xlsx)", type=["xlsx"])
with col2:
    uploaded_master = st.file_uploader("2️⃣ Upload File MASTER KATEGORI (.xlsx)", type=["xlsx"])

if uploaded_raw and uploaded_master:
    # Membaca kedua file
    df_raw = pd.read_excel(uploaded_raw, sheet_name=0)
    df_master = pd.read_excel(uploaded_master, sheet_name=0)
    
    # Cleansing awal: Hapus spasi tersembunyi di judul kolom
    df_raw.columns = df_raw.columns.str.strip()
    df_master.columns = df_master.columns.str.strip()
    
    # Validasi kolom wajib file mentahan
    required_raw = ['Tsh', 'Area', 'Name 1', 'Article Description', 'Quantity', 'Material Grp Desc. 2']
    missing_raw = [col for col in required_raw if col not in df_raw.columns]
    
    # Validasi kolom wajib file master
    required_master = ['Material', 'Category']
    missing_master = [col for col in required_master if col not in df_master.columns]
    
    if missing_raw:
        st.error(f"❌ File Mentahan kurang kolom: {', '.join(missing_raw)}. Pastikan kolom 'Tsh', 'Area', dan 'Material Grp Desc. 2' ada.")
    elif missing_master:
        st.error(f"❌ File Master kurang kolom: {', '.join(missing_master)}")
    else:
        st.success("✅ File berhasil dibaca. Sedang memproses penggabungan dan pembuatan Pivot...")
        
        # Cleansing Isi Data Raw
        for col in ['Tsh', 'Area', 'Name 1', 'Article Description', 'Material Grp Desc. 2']:
            df_raw[col] = df_raw[col].fillna('(kosong)').astype(str).str.strip()
            
        df_raw['Quantity'] = pd.to_numeric(df_raw['Quantity'], errors='coerce').fillna(0)
        
        # Cleansing Isi Data Master
        df_master['Material'] = df_master['Material'].astype(str).str.strip()
        df_master['Category'] = df_master['Category'].astype(str).str.strip()
        
        # 2. Proses VLOOKUP (Merge) berdasarkan Material
        df_merged = pd.merge(df_raw, df_master, left_on='Material Grp Desc. 2', right_on='Material', how='left')
        
        # Pisahkan menjadi 2 DataFrame (Device & Accessories)
        # Device = Kategori yang mengandung kata 'Device' atau 'Tablet'
        # Acc = Kategori yang mengandung kata 'Acc'
        df_device = df_merged[df_merged['Category'].str.contains('Device|Tablet', case=False, na=False)]
        df_acc = df_merged[df_merged['Category'].str.contains('Acc', case=False, na=False)]
        
        # 3. Fungsi Pembuat Pivot & Perapi Format Excel
        def generate_pivot_sheet(df_subset, sheet_name, excel_writer):
            if df_subset.empty:
                return # Lewati jika tidak ada datanya
                
            pivot_df = pd.pivot_table(
                df_subset, 
                index='Article Description', 
                columns=['Tsh', 'Area', 'Name 1'], 
                values='Quantity', 
                aggfunc='sum',
                fill_value=0
            )
            
            # Tambahkan kolom Total Keseluruhan
            pivot_df[('Total Keseluruhan', '', '')] = pivot_df.sum(axis=1)
            # Tambahkan baris Total Keseluruhan
            pivot_df.loc['Total Keseluruhan'] = pivot_df.sum(axis=0)
            
            # Simpan ke Sheet Excel
            pivot_df.to_excel(excel_writer, sheet_name=sheet_name)
            
            # Akses library openpyxl untuk merapikan format kolom bertingkat
            ws = excel_writer.sheets[sheet_name]
            
            # Penamaan sel pojok persis seperti SPR JABO
            ws['A1'] = 'Tsh'
            ws['A2'] = 'Area'
            ws['A3'] = 'Label Baris'
            
            # Hapus baris kosong ke-4
            ws.delete_rows(4)
            
            # Tambahkan Filter Dropdown di baris ke-3
            max_col = get_column_letter(ws.max_column)
            max_row = ws.max_row
            ws.auto_filter.ref = f"A3:{max_col}{max_row}"
            
            return pivot_df # Kembalikan dataframe untuk pratinjau web

        # 4. Tulis ke dalam File Excel (Memory Buffer)
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            pivot_dev = generate_pivot_sheet(df_device, 'Stock_Device', writer)
            pivot_acc = generate_pivot_sheet(df_acc, 'Stock_Accessories', writer)
            
        buffer.seek(0)
        
        # Tampilan Sekilas di Web (Opsional)
        st.info("💡 Pratinjau Data (Hanya menampilkan sebagian data Device. Silakan unduh untuk melihat format lengkap 2 Sheet).")
        if pivot_dev is not None:
            st.dataframe(pivot_dev.head(10), use_container_width=True)

        # 5. Tombol Unduh
        st.download_button(
            label="⬇️ Unduh Hasil Pivot SPR JABO (2 Sheet).xlsx",
            data=buffer.getvalue(),
            file_name="SPR_JABO_Device_&_Acc.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
elif uploaded_raw or uploaded_master:
    st.warning("⚠️ Mohon upload KEDUA file (Mentahan dan Master) untuk mulai memproses.")
